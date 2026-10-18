"""Exportación del reto: una fila por artículo, en el orden y con el News_ID del archivo original."""
from __future__ import annotations

import io
from datetime import timezone

import openpyxl
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from .db import utcnow
from .models import Analysis, Article, ImportBatch, Job

EXPORT_COLUMNS = ["News_ID", "Is_Risk", "Urgent", "Affected_Supplier_IDs"]
DANGEROUS_PREFIX = ("=", "+", "-", "@", "\t", "\r")


def latest_analysis(db: Session, article: Article) -> Analysis | None:
    return db.scalar(select(Analysis).where(Analysis.article_version_id == article.current_version_id)
                     .order_by(Analysis.created_at.desc(), Analysis.seq.desc()).limit(1))


def join_ids(ids) -> str:
    return "; ".join(sorted({i.strip() for i in ids if i and i.strip()}))


def batch_rows(db: Session, batch: ImportBatch) -> list[dict]:
    articles = list(db.scalars(select(Article).where(Article.batch_id == batch.id).order_by(Article.row_index)))
    # Un trabajo en cola o en curso para la versión vigente invalida cualquier resultado anterior.
    in_flight = {j.payload.get("version_id") for j in db.scalars(select(Job).where(Job.batch_id == batch.id, Job.status.in_(("queued", "running"))))}
    rows = []
    for art in articles:
        an = latest_analysis(db, art)
        an_pending = art.current_version_id in in_flight
        row = {"article_id": art.id, "news_id": art.external_id, "row_index": art.row_index, "analysis_id": an.id if an else None,
               "status": an.status if an else "not_started", "engine_mode": an.engine_mode if an else None,
               "is_risk": None, "urgent": None, "affected": [], "blocker": None, "review_reasons": [],
               "manual": False, "cached": bool(an and an.cached), "error_code": an.error_code if an else None,
               "duration_ms": an.duration_ms if an else None}
        # Solo en modo llm la segunda opinión cambia la decisión; en los demás no hay que esperarla.
        llm_pending = bool(an and (an.urgency_mode or "llm") == "llm" and any(a.llm_status == "pendiente" for a in an.assessments))
        if an is None or an.status in ("pending", "running") or an_pending or llm_pending:
            row["blocker"] = "pending"
            row["status"] = "pending" if an_pending else row["status"]
        elif an.engine_mode == "fixture":
            row["blocker"] = "fixture_result"
        elif an.status == "error":
            if an.manual_decision:
                md = an.manual_decision
                row.update(is_risk=md["is_risk"], urgent=md["urgent"], affected=md["affected_supplier_ids"], manual=True)
            else:
                row["blocker"] = "error"
        else:
            row.update(is_risk=an.is_risk, urgent=an.urgent, affected=an.affected_supplier_ids, review_reasons=an.review_reasons)
        if row["is_risk"] == "N":
            row["urgent"] = "N"
        rows.append(row)
    return rows


def export_status(db: Session, batch: ImportBatch) -> dict:
    rows = batch_rows(db, batch)
    blockers = [{"news_id": r["news_id"], "reason": r["blocker"], "error_code": r["error_code"]} for r in rows if r["blocker"]]
    return {"batch_id": batch.id, "rows": len(rows), "ready": not blockers and bool(rows), "blockers": blockers,
            "review_pending": sum(1 for r in rows if r["review_reasons"]), "manual": sum(1 for r in rows if r["manual"])}


def _text_cell(ws, row, col, value):
    cell = ws.cell(row=row, column=col)
    if value is None:
        return cell
    value = str(value)
    cell.value = value
    cell.data_type = "s"  # texto como texto: nunca fórmula
    return cell


def build_xlsx(db: Session, batch: ImportBatch, rows: list[dict] | None = None, extra_info: list[tuple] | None = None) -> bytes:
    """rows: por defecto las decisiones vigentes; scripts/export_consensus.py pasa las votadas por mayoría."""
    rows = rows if rows is not None else batch_rows(db, batch)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resultados"
    for c, h in enumerate(EXPORT_COLUMNS, 1):
        _text_cell(ws, 1, c, h).font = Font(bold=True)
    for r, row in enumerate(rows, 2):
        for c, v in enumerate([row["news_id"], row["is_risk"], row["urgent"], join_ids(row["affected"])], 1):
            _text_cell(ws, r, c, v)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["D"].width = 40

    # Copia de la hoja original con las columnas de entrada y las tres columnas de respuesta.
    copy = wb.create_sheet(batch.sheet_name[:31])
    input_cols = [c for c in batch.columns if c and c not in ("Event_Type", "Explanation")]
    for answer_col in ("Is_Risk", "Urgent", "Affected_Supplier_IDs"):
        if answer_col not in input_cols:
            input_cols.append(answer_col)
    for c, h in enumerate(input_cols, 1):
        _text_cell(copy, 1, c, h).font = Font(bold=True)
    articles = {a.id: a for a in db.scalars(select(Article).where(Article.batch_id == batch.id))}
    for r, row in enumerate(rows, 2):
        src = articles[row["article_id"]].source_row or {}
        for c, h in enumerate(input_cols, 1):
            if h == "Is_Risk":
                v = row["is_risk"]
            elif h == "Urgent":
                v = row["urgent"]
            elif h == "Affected_Supplier_IDs":
                v = join_ids(row["affected"])
            else:
                v = src.get(h)
            _text_cell(copy, r, c, v)

    report = wb.create_sheet("Informe")
    analyses = [an for an in (db.get(Analysis, r["analysis_id"]) for r in rows if r.get("analysis_id")) if an is not None]
    info = [
        ("Lote", batch.id), ("Archivo", batch.file_name), ("Hoja", batch.sheet_name), ("SHA-256", batch.file_sha256),
        ("Generado (UTC)", utcnow().astimezone(timezone.utc).isoformat()),
        ("Modelos", ", ".join(sorted({a.model_id or "-" for a in analyses}))),
        ("Política", ", ".join(sorted({a.policy_version for a in analyses}))),
        ("Versión de preguntas", ", ".join(sorted({a.question_version for a in analyses}))),
        ("Filas", len(rows)), ("Con revisión pendiente", sum(1 for r in rows if r["review_reasons"])),
        ("Decisión manual", sum(1 for r in rows if r["manual"])), ("Resultados en caché", sum(1 for r in rows if r["cached"])),
        ("Nota", "Prioridad 0–100 = orden de atención, no probabilidad. Revisiones pendientes exportan su mejor decisión binaria."),
        *(extra_info or []),
    ]
    info = [(k, v) for k, v in info if v not in (None, "")]
    for r, (k, v) in enumerate(info, 1):
        _text_cell(report, r, 1, k).font = Font(bold=True)
        _text_cell(report, r, 2, v)
    start = len(info) + 2
    for c, h in enumerate(["News_ID", "Estado", "Motivos de revisión", "Decisión manual", "Caché", "Duración ms"], 1):
        _text_cell(report, start, c, h).font = Font(bold=True)
    for r, row in enumerate(rows, start + 1):
        for c, v in enumerate([row["news_id"], row["status"], ", ".join(row["review_reasons"]), "sí" if row["manual"] else "",
                               "sí" if row["cached"] else "", row["duration_ms"]], 1):
            _text_cell(report, r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def csv_safe(value: str) -> str:
    """Neutraliza fórmulas al generar CSV."""
    return "'" + value if value and value.startswith(DANGEROUS_PREFIX) else value
